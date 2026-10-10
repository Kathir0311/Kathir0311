"""Excel report for an accessibility run."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .model import DocModel, Finding

FILL = {
    "ERROR": PatternFill("solid", fgColor="FDDCDC"),
    "WARNING": PatternFill("solid", fgColor="FFF1CC"),
    "PASS": PatternFill("solid", fgColor="DDF2E5"),
    "INFO": PatternFill("solid", fgColor="E8EEF6"),
}
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="273947")
WRAP = Alignment(wrap_text=True, vertical="top")


def _cell(v):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        return v[:32000]
    return v


def _append(ws, row) -> None:
    """Worksheet.append that never fails on control characters in PDF text."""
    ws.append([_cell(v) for v in row])


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    _append(ws, headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.font = HEAD
        cell.fill = HEAD_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    return ws


def _findings(wb, title, findings: list[Finding]):
    ws = _sheet(wb, title, ["Code", "Severity", "Checkpoint", "Clause", "Detail", "Pages",
                            "Count", "Fix type"], [11, 11, 46, 20, 64, 22, 8, 10])
    for f in findings:
        _append(ws, f.as_row())
        fill = FILL.get(f.severity.upper())
        for col in range(1, 9):
            cell = ws.cell(row=ws.max_row, column=col)
            if fill:
                cell.fill = fill
            cell.alignment = WRAP


def write_report(path: str, *, filename: str, before: list[Finding], after: list[Finding],
                 before_score: dict, after_score: dict, model: DocModel | None = None,
                 build_stats: dict | None = None, preflight: dict | None = None,
                 options: dict | None = None, verapdf: dict | None = None,
                 contrast: dict | None = None) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _append(ws, ["Dezquare PDF Accessibility Studio — accessibility report"])
    ws["A1"].font = Font(bold=True, size=14)
    _append(ws, [])
    bs = build_stats or {}
    rows = [
        ("File", filename),
        ("Pages", model.page_count if model else ""),
        ("Tagging mode", (model.mode if model else "")),
        ("Title", model.title if model else ""),
        ("Language", model.lang if model else ""),
        ("Checks run", after_score.get("checked", 0)),
        ("Errors  before → after", f"{before_score.get('errors', 0)} → {after_score.get('errors', 0)}"),
        ("Warnings  before → after", f"{before_score.get('warnings', 0)} → {after_score.get('warnings', 0)}"),
        ("Pass rate  before → after", f"{before_score.get('percent', 0)}% → {after_score.get('percent', 0)}%"),
        ("veraPDF", (verapdf or {}).get("verdict", "not run")),
        ("Structure elements written", bs.get("elements", "")),
        ("Marked-content IDs", bs.get("mcids", "")),
        ("Text operators split at tag boundaries", bs.get("split_ops", "")),
        ("Content outside every tag", f"{bs.get('unmatched', 0)} glyph(s), "
                                     f"{bs.get('attached_to_nearest', 0)} attached to nearest tag"),
        ("Annotations tagged", f"{bs.get('annots_tagged', 0)} (+{bs.get('annots_fallback', 0)} placed automatically)"),
        ("Empty tags removed", len(bs.get("pruned", []) or [])),
        ("Figures still without alt text", len(bs.get("missing_alt", []) or [])),
        ("Optional-content markers kept", bs.get("kept_oc", 0)),
    ]
    if contrast:
        iss = contrast.get("issues", [])
        fixed = sum(1 for i in iss if i.get("status") == "fixed")
        rows.append((f"Colour contrast ({contrast.get('level', 'AA')})",
                     f"{contrast.get('checked', 0)} text runs measured · {len(iss)} below the minimum · "
                     f"{fixed} fixed · {len(iss) - fixed} open"))
    for k, v in rows:
        _append(ws, [k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 90

    def section(title, items):
        if not items:
            return
        _append(ws, [])
        _append(ws, [title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for it in items:
            _append(ws, ["", it])
            ws.cell(row=ws.max_row, column=2).alignment = WRAP

    section("Preflight", (preflight or {}).get("log", []))
    section("Preflight warnings", (preflight or {}).get("warnings", []))
    section("Automatic fixes", bs.get("applied", []))
    section("Needs a human check", bs.get("review", []) + (model.warnings if model else []))
    _append(ws, [])
    _append(ws, ["", "Machine checks cover what software can decide. Confirm final conformance "
                   "with PAC 2024 and/or veraPDF, and work through the Manual checklist sheet."])
    ws.cell(row=ws.max_row, column=2).alignment = WRAP

    _findings(wb, "Findings after", after)
    _findings(wb, "Findings before", before)
    if verapdf and verapdf.get("findings"):
        _findings(wb, "veraPDF", verapdf["findings"])

    if model is not None:
        depth = model.depth_of()
        ts = _sheet(wb, "Tag tree", ["#", "Page", "Depth", "Tag", "Text / alt", "Lang",
                                     "Confidence", "Source", "Why"],
                    [7, 6, 6, 12, 80, 7, 11, 9, 46])
        for i, n in enumerate(model.ordered_nodes(), start=1):
            label = n.alt if n.role in ("Figure", "Formula") and n.alt else n.text
            _append(ts, [i, n.page + 1 if n.page >= 0 else "", depth.get(n.id, 0),
                       "Artifact" if n.artifact else n.role,
                       ("    " * depth.get(n.id, 0)) + (label or "")[:300], n.lang,
                       round(n.confidence, 2), n.source, n.note])
            if n.confidence < 0.55 and not n.artifact:
                for col in range(1, 10):
                    ts.cell(row=ts.max_row, column=col).fill = FILL["WARNING"]

        fs = _sheet(wb, "Figures", ["Page", "Tag", "Alternative text", "Source", "Status"],
                    [7, 10, 90, 10, 14])
        for n in model.ordered_nodes():
            if n.role in ("Figure", "Formula") and not n.artifact:
                status = "OK" if (n.alt or n.actual_text) else "MISSING"
                _append(fs, [n.page + 1, n.role, n.alt or n.actual_text, n.source, status])
                if status == "MISSING":
                    fs.cell(row=fs.max_row, column=5).fill = FILL["ERROR"]

        tb = _sheet(wb, "Tables", ["Page", "Rows", "Columns", "Header cells", "Merged cells",
                                   "Header association", "Note"], [7, 7, 9, 13, 13, 20, 60])
        for n in model.ordered_nodes():
            if n.role != "Table" or n.artifact:
                continue
            cells = [d for d in model.descendants(n.id) if d.role in ("TH", "TD")]
            rows = [d for d in model.descendants(n.id) if d.role == "TR"]
            cols = max((sum(model.nodes[c].colspan for c in r.children
                            if model.nodes[c].role in ("TH", "TD")) for r in rows), default=0)
            ths = sum(1 for d in cells if d.role == "TH")
            merged = sum(1 for d in cells if d.rowspan > 1 or d.colspan > 1)
            assoc = "Headers/IDs" if any(d.headers for d in cells) else (
                "Scope" if any(d.scope for d in cells) else "none")
            _append(tb, [n.page + 1, len(rows), cols, ths, merged, assoc, n.note])

        ls = _sheet(wb, "Links & fields", ["Page", "Tag", "Text / label", "Annotation"],
                    [7, 10, 80, 14])
        for n in model.ordered_nodes():
            if n.role in ("Link", "Form", "Annot"):
                _append(ls, [n.page + 1, n.role, n.text, ", ".join(str(a) for a in n.annots)])

    if contrast and contrast.get("issues"):
        cs = _sheet(wb, "Contrast", ["Page", "Text", "Size", "Large", "Text colour", "Background",
                                     "Ratio", "Required", "Where", "Fix (text, box or outline)", "Status",
                                     "Ratio after", "Note"],
                    [7, 48, 7, 7, 12, 12, 8, 9, 10, 15, 13, 11, 60])
        for i in contrast["issues"]:
            how = i.get("fixed_with") or ("background" if i.get("suggested_bg") else "text")
            fix_txt = {"stroke": i.get("fixed_detail") or "outline",
                       "background": "box " + (i.get("suggested_bg") or ""),
                       "text": i.get("suggested") or ""}.get(how, "")
            _append(cs, [i["page"] + 1, i["text"], i["size"], "yes" if i["large"] else "",
                         i["fg"], i["bg"], i["ratio"], i["required"], i["zone"], fix_txt,
                         i.get("status", "open"), i.get("ratio_after", ""),
                         i.get("fix_note") or i.get("note", "")])
            r = cs.max_row
            swatch = {"background": "suggested_bg", "text": "suggested"}.get(how, "")
            for col, key in ((5, "fg"), (6, "bg")) + (((10, swatch),) if swatch else ()):
                hx = (i.get(key) or "").lstrip("#")
                if len(hx) == 6:
                    cs.cell(row=r, column=col).fill = PatternFill("solid", fgColor=hx)
                    lum = (int(hx[0:2], 16) * 299 + int(hx[2:4], 16) * 587 + int(hx[4:6], 16) * 114) / 1000
                    cs.cell(row=r, column=col).font = Font(color="000000" if lum > 140 else "FFFFFF")
            st = i.get("status", "open")
            cs.cell(row=r, column=11).fill = FILL["PASS"] if st == "fixed" else FILL["ERROR"] if i["severity"] == "error" else FILL["WARNING"]

    if bs.get("pruned"):
        ps = _sheet(wb, "Removed empty tags", ["Page", "Tag", "Text"], [7, 12, 90])
        for p in bs["pruned"]:
            _append(ps, [p.get("page", -1) + 1, p.get("role"), p.get("text")])

    if bs.get("unmatched_pages"):
        us = _sheet(wb, "Unmatched content", ["Page", "Glyphs outside every tag"], [8, 26])
        for p, n in sorted(bs["unmatched_pages"].items(), key=lambda t: int(t[0])):
            _append(us, [int(p) + 1, n])

    ms = _sheet(wb, "Manual checklist", ["Checkpoint", "How to verify", "Status", "Notes"],
                [46, 70, 12, 40])
    checklist = [
        ("Reading order", "Open the reading-order view; read the page numbers top to bottom."),
        ("Heading levels reflect the outline", "Compare the Tag tree sheet with the visual hierarchy."),
        ("Alternative text is accurate", "Read each entry on the Figures sheet against the image."),
        ("Decorative images are artifacts", "Borders, dividers, background textures, repeated logos."),
        ("Tables: headers correct", "Every data cell announces the right row/column headers."),
        ("Lists are real lists", "Bullets/numbers are in Lbl, text in LBody, nesting correct."),
        ("Link text is meaningful", "Avoid 'click here'; the Links sheet shows each link's text."),
        ("Form fields have tooltips", "Tooltips should match visible labels."),
        ("Colour contrast ≥ 4.5:1", "Check text on coloured backgrounds with a contrast tool."),
        ("Colour is not the only cue", "Charts and required fields have a non-colour indicator."),
        ("Language of foreign passages", "Passages in another language carry a Lang value."),
        ("Final validation", "Run PAC 2024 and/or veraPDF on the delivered file."),
    ]
    dv = DataValidation(type="list", formula1='"Pass,Fail,N/A"', allow_blank=True)
    ms.add_data_validation(dv)
    for t, h in checklist:
        _append(ms, [t, h, "", ""])
        dv.add(ms.cell(row=ms.max_row, column=3))
        ms.cell(row=ms.max_row, column=2).alignment = WRAP

    if options:
        os_ = _sheet(wb, "Options used", ["Option", "Value"], [34, 30])
        for k, v in options.items():
            _append(os_, [k, str(v)])

    wb.save(path)
    return path


def write_batch_summary(path: str, rows: list[dict]) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch"
    headers = ["File", "Pages", "Mode", "Errors before", "Errors after", "Warnings after",
               "Pass rate after", "Figures without alt", "Contrast issues", "veraPDF", "Status", "Output"]
    _append(ws, headers)
    for i, w in enumerate([40, 7, 9, 13, 12, 14, 14, 18, 15, 30, 12, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        ws.cell(row=1, column=i).font = HEAD
        ws.cell(row=1, column=i).fill = HEAD_FILL
    for r in rows:
        _append(ws, [r.get(h.lower().replace(" ", "_"), "") for h in headers])
        status = str(r.get("status", "")).lower()
        if status == "error":
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = FILL["ERROR"]
    ws.freeze_panes = "A2"
    wb.save(path)
    return path
